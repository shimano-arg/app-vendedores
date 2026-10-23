"""Carga semanal de OITM.AvgPrice (planilla Juan Paiva) a costo_snapshot_semanal.

Lee un Excel exportado de la consulta AVG PRICE de SAP con columnas
ItemCode y AvgPrice (tolera encabezados en espanol o con espacios: 'Item
Code', 'Codigo', 'Codigo Articulo', 'Precio Promedio', 'Avg Price', etc.)
y lo carga a `app-vendedores-shimano.shimano_app.costo_snapshot_semanal`
con:
    snapshot_week_start = lunes de la semana de la fecha pasada por parametro.
    snapshot_at         = CURRENT_TIMESTAMP().
    cost_source         = 'paiva_import_semanal_<yyyymmdd>'  (yyyymmdd = fecha
                          cruda del parametro, NO el lunes; para auditar
                          cuando se cargo cada semana).

Idempotente: DELETE + INSERT por snapshot_week_start (si se corre dos veces
el mismo lunes, pisa).

Guards (abortan sin cargar):
  * < 500 filas con AvgPrice > 0.
  * ItemCode duplicados en el Excel (misma clave con 2+ precios distintos).

Autenticacion: Application Default Credentials (ADC) local. Correr
`gcloud auth application-default login` una vez por maquina. No requiere
archivo SA key.

Uso:
    python scripts/load_paiva_semanal.py --excel path/a/avg_price.xlsx --fecha 2026-10-27
    python scripts/load_paiva_semanal.py --excel ... --fecha ... --dry-run

Con --dry-run imprime el SQL DELETE + INSERT que se ejecutaria + el conteo
de filas, sin tocar BQ.
"""
import argparse
import re
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

try:
    import openpyxl  # noqa: F401
except ImportError:
    print('[FATAL] falta openpyxl. Instalar: pip install openpyxl', file=sys.stderr)
    sys.exit(2)

from google.cloud import bigquery


BQ_PROJECT = 'app-vendedores-shimano'
BQ_DATASET = 'shimano_app'
BQ_LOCATION = 'southamerica-east1'
TABLE = f'{BQ_PROJECT}.{BQ_DATASET}.costo_snapshot_semanal'

MIN_FILAS_VALIDAS = 500


def _norm_header(s: str) -> str:
    """Normaliza encabezado para matching tolerante (lower + sin espacios ni acentos)."""
    s = (s or '').strip().lower()
    s = s.replace(' ', '').replace('_', '').replace('-', '')
    s = s.replace('á', 'a').replace('é', 'e').replace('í', 'i').replace('ó', 'o').replace('ú', 'u').replace('ñ', 'n')
    return s


# Candidatos de encabezados aceptados (normalizados).
ITEM_CODE_HEADERS = {
    'itemcode', 'codigo', 'codigoarticulo', 'codigoitem',
    'articulo', 'sku', 'cod',
}
AVG_PRICE_HEADERS = {
    'avgprice', 'avgstdprice', 'preciopromedio', 'precioavg', 'avg',
    'costo', 'costopromedio', 'costoavg', 'costounitario',
}


def lunes_de_la_semana(d: date) -> date:
    """Lunes de la semana ISO de la fecha dada (si d es lunes, devuelve d)."""
    return d - timedelta(days=d.weekday())


def read_excel(path: Path) -> list[tuple[str, float]]:
    """Lee el Excel y devuelve [(item_code, avg_price), ...].

    Toma la primera hoja. Localiza columnas por header normalizado tolerando
    espanol y espacios. Lanza SystemExit si no las encuentra.
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        print(f'[FATAL] Excel vacio: {path}', file=sys.stderr)
        sys.exit(3)

    norm_to_idx = {}
    for idx, h in enumerate(header_row):
        if h is None:
            continue
        norm_to_idx[_norm_header(str(h))] = idx

    item_idx = next((norm_to_idx[k] for k in ITEM_CODE_HEADERS if k in norm_to_idx), None)
    price_idx = next((norm_to_idx[k] for k in AVG_PRICE_HEADERS if k in norm_to_idx), None)

    if item_idx is None or price_idx is None:
        print(f'[FATAL] no encuentro columnas ItemCode y AvgPrice. '
              f'Headers vistos (normalizados): {sorted(norm_to_idx.keys())}',
              file=sys.stderr)
        print(f'  Esperaba una de ItemCode: {sorted(ITEM_CODE_HEADERS)}', file=sys.stderr)
        print(f'  Esperaba una de AvgPrice: {sorted(AVG_PRICE_HEADERS)}', file=sys.stderr)
        sys.exit(4)

    out: list[tuple[str, float]] = []
    for row in rows_iter:
        if not row:
            continue
        item = row[item_idx] if item_idx < len(row) else None
        price = row[price_idx] if price_idx < len(row) else None
        if item is None:
            continue
        item_str = str(item).strip()
        if not item_str:
            continue
        try:
            price_f = float(price) if price is not None else 0.0
        except (TypeError, ValueError):
            continue
        out.append((item_str, price_f))
    return out


def validate(data: list[tuple[str, float]]) -> dict:
    """Guards: < 500 con AvgPrice>0 abort; ItemCode duplicados abort."""
    con_precio = [(i, p) for i, p in data if p > 0]
    if len(con_precio) < MIN_FILAS_VALIDAS:
        print(f'[FATAL] solo {len(con_precio)} filas con AvgPrice > 0 '
              f'(minimo {MIN_FILAS_VALIDAS}). Abortando sin cargar.',
              file=sys.stderr)
        sys.exit(5)
    seen: dict[str, float] = {}
    dupes: list[tuple[str, float, float]] = []
    for item, price in con_precio:
        if item in seen and abs(seen[item] - price) > 0.01:
            dupes.append((item, seen[item], price))
        seen[item] = price
    if dupes:
        print(f'[FATAL] ItemCode duplicados en el Excel con precios distintos '
              f'({len(dupes)} casos). Primeros 5:', file=sys.stderr)
        for item, p1, p2 in dupes[:5]:
            print(f'    {item}: {p1} vs {p2}', file=sys.stderr)
        sys.exit(6)
    return {
        'total': len(data),
        'con_precio': len(con_precio),
        'unicos': len(seen),
    }


def build_sql(data: list[tuple[str, float]], week_start: date, cost_source: str) -> list[str]:
    """Devuelve 2 statements SQL: DELETE idempotente + INSERT batch.

    Usa literales de BQ (STRUCT) para evitar query parameters en multi-stmt
    scripts (BQ no acepta parametros en todas las statements de un script).
    """
    # DELETE idempotente del partition week_start.
    sql_delete = (
        f"DELETE FROM `{TABLE}` "
        f"WHERE snapshot_week_start = DATE('{week_start.isoformat()}');"
    )
    # INSERT batch via UNNEST(STRUCT).
    pairs = []
    for item, price in data:
        if price <= 0:
            continue
        safe_item = item.replace("'", "''")
        pairs.append(f"STRUCT('{safe_item}' AS item_code, {price} AS cost_avg_ars)")
    values_sql = ', '.join(pairs)
    sql_insert = f"""
INSERT INTO `{TABLE}`
  (snapshot_week_start, snapshot_at, item_code, cost_avg_ars, cost_source)
SELECT
  DATE('{week_start.isoformat()}'),
  CURRENT_TIMESTAMP(),
  t.item_code,
  t.cost_avg_ars,
  '{cost_source}'
FROM UNNEST([{values_sql}]) AS t;
""".strip()
    return [sql_delete, sql_insert]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--excel', required=True, help='Path al Excel exportado de SAP.')
    ap.add_argument('--fecha', required=True,
                    help='Fecha ISO (YYYY-MM-DD) para calcular lunes de la semana + cost_source.')
    ap.add_argument('--dry-run', action='store_true',
                    help='No ejecuta BQ, imprime SQL + conteos.')
    args = ap.parse_args()

    excel_path = Path(args.excel)
    if not excel_path.exists():
        print(f'[FATAL] no existe: {excel_path}', file=sys.stderr)
        sys.exit(1)

    try:
        fecha = date.fromisoformat(args.fecha)
    except ValueError:
        print(f'[FATAL] --fecha invalida: {args.fecha} (esperado YYYY-MM-DD)',
              file=sys.stderr)
        sys.exit(1)

    week_start = lunes_de_la_semana(fecha)
    yyyymmdd = fecha.strftime('%Y%m%d')
    cost_source = f'paiva_import_semanal_{yyyymmdd}'

    print(f'[input]  excel         = {excel_path}')
    print(f'[input]  fecha         = {fecha.isoformat()}')
    print(f'[calc]   week_start    = {week_start.isoformat()}  (lunes de la semana)')
    print(f'[calc]   cost_source   = {cost_source}')
    print()

    data = read_excel(excel_path)
    print(f'[read]   filas         = {len(data)}')

    stats = validate(data)
    print(f'[check]  con AvgPrice>0 = {stats["con_precio"]}')
    print(f'[check]  items unicos   = {stats["unicos"]}')

    stmts = build_sql(data, week_start, cost_source)
    print()
    if args.dry_run:
        print('[DRY_RUN] NO ejecuta BQ. SQL que se ejecutaria:')
        print()
        print('--- DELETE ---')
        print(stmts[0])
        print()
        print('--- INSERT (primeros 500 chars) ---')
        print(stmts[1][:500] + ('...[truncado]...' if len(stmts[1]) > 500 else ''))
        return

    print(f'[exec]   target        = {TABLE}')
    c = bigquery.Client(project=BQ_PROJECT, location=BQ_LOCATION)
    # Ejecutar como script multi-statement.
    full_sql = stmts[0] + '\n' + stmts[1]
    job = c.query(full_sql)
    job.result()
    print(f'[exec]   OK bytes_processed={job.total_bytes_processed or 0:,}')

    # Verif post.
    q = f"""
    SELECT COUNT(*) AS n, COUNT(DISTINCT item_code) AS items,
      ROUND(AVG(cost_avg_ars), 2) AS avg,
      ROUND(MIN(cost_avg_ars), 2) AS min,
      ROUND(MAX(cost_avg_ars), 2) AS max
    FROM `{TABLE}`
    WHERE snapshot_week_start = DATE('{week_start.isoformat()}')
      AND cost_source = '{cost_source}'
    """
    r = list(c.query(q).result())[0]
    print(f'[check]  post rows      = {r.n}  items_unique = {r.items}  '
          f'avg = {r.avg}  min = {r.min}  max = {r.max}')


if __name__ == '__main__':
    main()
